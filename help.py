from openpyxl import load_workbook
import re

# Učitaj fajl
wb = load_workbook("Konekcije.xlsx")
ws = wb.active

# Regex da prepoznaš lokacije (velika slova, zagrade, brojevi)
lokacija_re = re.compile(r"^[A-ZČĆŽŠĐ0-9() /.-]{2,}$")

rezultat = []

# 1️⃣ Prođi svaku kolonu posebno
for col_idx in range(1, ws.max_column + 1):
    lokacija = None
    lokacija_row = None

    for row_idx in range(1, ws.max_row + 1):
        cell_value = ws.cell(row=row_idx, column=col_idx).value
        if cell_value is None:
            continue

        cell_text = str(cell_value).strip()

        # Ako je ćelija lokacija
        if lokacija_re.match(cell_text):
            lokacija = cell_text
            lokacija_row = row_idx
            continue

        # Ako već imamo aktivnu lokaciju — sve ispod pripada njoj
        if lokacija and row_idx > lokacija_row:
            rezultat.append({
                "lokacija": lokacija,
                "red": row_idx,
                "kolona": col_idx,
                "vrednost": cell_text
            })

# 2️⃣ Prikaz rezultata
for r in rezultat[:20]:
    print(r)

print(f"\nUkupno redova: {len(rezultat)}")
