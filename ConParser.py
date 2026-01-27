from openpyxl import load_workbook
import os
import re


class ConParser:
    def __init__(self):
        self.file = os.path.join("Konekcije.xlsx")

    def find_connection(self, label):
        try:
            wb = load_workbook(self.file, read_only=True)
            sheet = wb.active

            # for i, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            #     red_str = " ".join(map(str, filter(None, row)))

            #     if label.lower() in red_str.lower():
            #         print(f"Nađen u redu {i}: {row}")
            #     pattern = re.compile(r'(unit-[^\s]+.*?)(?=unit-|$)', re.DOTALL)
            pattern = re.compile(
                r'(unit-[^\s]+\s+port-[^\s]+.*?)(?=unit-[^\s]+\s+port-[^\s]+|$)', re.DOTALL)
            for i, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                red_str = " ".join(map(str, filter(None, row)))
                blokovi = pattern.findall(red_str)
                for blok in blokovi:
                    if label.lower().replace(" ", "") in blok.lower().replace(" ", ""):
                        return blok.strip()
            return None
        except Exception as e:
            print(e)

    def find_next_mux(self, connection, label):
        try:
            parts = connection.split(label, 1)
            pattern = r'^unit-\d+\s+port-\d+$'
            if re.fullmatch(pattern, parts[0].strip()):
                text = parts[1].strip()
            else:
                text = parts[0].strip()
            pattern = r'^(?:(?P<unit>unit-\d+)\s+(?P<port>port-\d+)\s+(?P<label>.+)|(?P<label2>.+?)\s+(?P<port2>port-\d+)\s+(?P<unit2>unit-\d+))'

            m = re.match(pattern, text)
            if m:
                if m.group('unit'):
                    parts = [m.group('label'), m.group(
                        'unit'), m.group('port')]
                else:
                    parts = [m.group('label2'), m.group(
                        'unit2'), m.group('port2')]
                return parts
            else:
                return None
        except Exception as e:
            print(e)


if __name__ == "__main__":
    parser = ConParser()
    # parser.find_connection("port-1: LL 10")
    parser.find_next_mux(
        "unit-16 port-1 LL 10  port-1: LL 10 port-1 unit-16", "port-1: LL 10")
